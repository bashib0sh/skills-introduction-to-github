"""Build a sample tranched (senior + mezzanine) debt sizing model for a solar PV project."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

OUT = "/home/user/skills-introduction-to-github/models/tranched_debt_sizing.xlsx"
N_YEARS = 25
FIRST, LAST = 4, 4 + N_YEARS  # col D = year 0 (financial close), E..AC = operating years 1..25
FC = L(FIRST)
C1, CN = L(FIRST + 1), L(LAST)

BLUE = Font(name="Arial", size=10, color="0000FF")
GREEN = Font(name="Arial", size=10, color="008000")
BLACK = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=14, bold=True)
HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
TOP = Border(top=Side(style="thin"))

MM = '#,##0.0;(#,##0.0);"-"'
INT = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
X = '0.00"x";(0.00"x");"-"'
YR = '0'

wb = Workbook()

# ---------------------------------------------------------------- Inputs
inp = wb.active
inp.title = "Inputs"
inp["A1"] = "Inputs: 150 MW Solar PV Project (Illustrative)"
inp["A1"].font = TITLE
inp["A2"] = ("Legend: blue = hardcoded input (edit these) | yellow fill = key sizing levers | "
             "black = formula. All values are illustrative placeholders, not market quotes.")
inp["A2"].font = Font(name="Arial", size=9, italic=True)

rows = [
    ("Project", None),
    ("Installed capacity", 150, "MW", INT, "Illustrative"),
    ("P50 net capacity factor", 0.24, "%", PCT, "Illustrative; replace with independent energy yield report"),
    ("P90 (1-year) net capacity factor", 0.215, "%", PCT, "Illustrative; replace with independent energy yield report"),
    ("Annual module degradation", 0.005, "%", PCT, "Illustrative"),
    ("Operating life", 25, "years", INT, "Model horizon from COD"),
    ("Revenue", None),
    ("PPA tariff (Year 1)", 68, "$/MWh", '#,##0.00', "Illustrative"),
    ("Tariff escalation", 0.02, "% p.a.", PCT, "Illustrative"),
    ("Operating costs", None),
    ("Fixed O&M + SG&A (Year 1)", 18, "$/kW-yr", '#,##0.00', "Illustrative"),
    ("Opex escalation", 0.025, "% p.a.", PCT, "Illustrative"),
    ("Project cost", None),
    ("EPC + development costs", 165, "$mm", MM, "Illustrative"),
    ("Financing fees + IDC (estimate)", 10, "$mm", MM, "Held fixed to avoid a circular reference; true-up in a full model"),
    ("Total project cost", "=B17+B18", "$mm", MM, "Formula"),
    ("Tax", None),
    ("Corporate tax rate", 0.25, "%", PCT, "Illustrative"),
    ("Tax depreciation life (straight line)", 20, "years", INT, "Interest tax shield excluded (conservative, avoids circularity)"),
    ("Senior term loan (Tranche A)", None),
    ("Tenor from COD", 18, "years", INT, "Illustrative"),
    ("All-in interest rate (base + margin + swap)", 0.065, "% p.a.", PCT, "Illustrative; fixed via swap"),
    ("Target DSCR: P50 case", 1.30, "x", X, "Senior sculpting target, P50"),
    ("Target DSCR: P90 case", 1.20, "x", X, "Senior sculpting target, P90 (1-yr)"),
    ("Max senior gearing", 0.80, "% of cost", PCT, "Senior debt / total project cost cap"),
    ("Mezzanine / HoldCo loan (Tranche B)", None),
    ("Tenor from COD", 18, "years", INT, "Coterminous with senior; a longer tenor pushes repayment into the post-senior tail (see Summary note 3)"),
    ("All-in interest rate", 0.10, "% p.a.", PCT, "Illustrative"),
    ("Combined DSCR floor: P50 case", 1.15, "x", X, "Senior + mezz debt service"),
    ("Combined DSCR floor: P90 case", 1.05, "x", X, "Senior + mezz debt service"),
    ("Max total gearing", 0.85, "% of cost", PCT, "(Senior + mezz) / total project cost cap"),
]

inp["A3"], inp["B3"], inp["C3"], inp["D3"] = "Item", "Value", "Unit", "Source / note"
for c in "ABCD":
    inp[f"{c}3"].font = HDR
    inp[f"{c}3"].fill = HDR_FILL
r = 4
for row in rows:
    if row[1] is None:
        inp.cell(r, 1, row[0]).font = BOLD
        for c in range(1, 5):
            inp.cell(r, c).fill = SEC_FILL
    else:
        label, val, unit, fmt, note = row
        inp.cell(r, 1, label).font = BLACK
        cell = inp.cell(r, 2, val)
        cell.font = BLACK if isinstance(val, str) else BLUE
        cell.number_format = fmt
        inp.cell(r, 3, unit).font = BLACK
        inp.cell(r, 4, note).font = Font(name="Arial", size=9, italic=True, color="595959")
    r += 1

# Input cell map (rows as laid out above)
I = dict(cap="Inputs!$B$5", cf50="Inputs!$B$6", cf90="Inputs!$B$7", deg="Inputs!$B$8",
         life="Inputs!$B$9", tariff="Inputs!$B$11", tesc="Inputs!$B$12", opex="Inputs!$B$14",
         oesc="Inputs!$B$15", epc="Inputs!$B$17", fees="Inputs!$B$18", cost="Inputs!$B$19",
         tax="Inputs!$B$21", dep="Inputs!$B$22",
         s_ten="Inputs!$B$24", s_rate="Inputs!$B$25", s_d50="Inputs!$B$26", s_d90="Inputs!$B$27",
         s_gear="Inputs!$B$28",
         m_ten="Inputs!$B$30", m_rate="Inputs!$B$31", m_d50="Inputs!$B$32", m_d90="Inputs!$B$33",
         m_gear="Inputs!$B$34")
for k in ("s_d50", "s_d90", "s_gear", "m_d50", "m_d90", "m_gear"):
    inp[I[k].split("!")[1].replace("$", "")].fill = KEY_FILL
inp.column_dimensions["A"].width = 44
inp.column_dimensions["B"].width = 12
inp.column_dimensions["C"].width = 11
inp.column_dimensions["D"].width = 70
inp.freeze_panes = "A4"

# ---------------------------------------------------------------- Model
m = wb.create_sheet("Model")
m["A1"] = "Tranched Debt Sizing Model: Annual ($mm unless stated)"
m["A1"].font = TITLE
m["A2"] = ("Sizing: each tranche is sculpted to the LOWER of its P50 and P90 DSCR targets, "
           "then capped by gearing. Mezz sizes on combined DSCR and takes the residual capacity after senior.")
m["A2"].font = Font(name="Arial", size=9, italic=True)

ROW = {}


def section(r, title):
    for c in range(1, LAST + 1):
        m.cell(r, c).fill = SEC_FILL
    m.cell(r, 1, title).font = BOLD


def line(r, key, label, unit, fmt, first=None, each=None, total=None, font=BLACK, bold=False):
    """first: formula for year-0 col (or None); each: callable(col_letter, prev_letter) for years 1..N."""
    ROW[key] = r
    m.cell(r, 1, label).font = BOLD if bold else BLACK
    m.cell(r, 2, unit).font = BLACK
    if total is not None:
        c = m.cell(r, 3, total)
        c.font = BOLD if bold else font
        c.number_format = fmt
    if first is not None:
        c = m.cell(r, FIRST, first)
        c.font = font
        c.number_format = fmt
    if each is not None:
        for ci in range(FIRST + 1, LAST + 1):
            c = m.cell(r, ci, each(L(ci), L(ci - 1)))
            c.font = font
            c.number_format = fmt
            if bold:
                c.font = BOLD
    if bold:
        for ci in range(3, LAST + 1):
            m.cell(r, ci).border = TOP


# Header / timeline
for ci, h in [(1, "Line item"), (2, "Unit"), (3, "Total / value")]:
    m.cell(4, ci, h)
m.cell(4, FIRST, 0)
for ci in range(FIRST + 1, LAST + 1):
    m.cell(4, ci, f"={L(ci - 1)}4+1")
for ci in range(1, LAST + 1):
    m.cell(4, ci).font = HDR
    m.cell(4, ci).fill = HDR_FILL
    m.cell(4, ci).alignment = Alignment(horizontal="center" if ci > 2 else "left")
    if ci >= FIRST:
        m.cell(4, ci).number_format = '"Year "0'
m.cell(4, FIRST).font = Font(name="Arial", size=10, bold=True, color="FFFF00")
m.cell(4, FIRST).comment = Comment("Year 0 = financial close / funding date. Years 1-25 = operations from COD.", "Model")
ROW["yr"] = 4

line(5, "op", "Operations flag", "flag", "0", each=lambda c, p: f"=IF(AND({c}$4>=1,{c}$4<={I['life']}),1,0)")
line(6, "sf", "Senior repayment flag", "flag", "0", each=lambda c, p: f"=IF(AND({c}$4>=1,{c}$4<={I['s_ten']}),1,0)")
line(7, "mf", "Mezz repayment flag", "flag", "0", each=lambda c, p: f"=IF(AND({c}$4>=1,{c}$4<={I['m_ten']}),1,0)")

section(9, "Operations & CFADS")
gen = lambda cf: (lambda c, p: f"={I['cap']}*{I[cf]}*8760*(1-{I['deg']})^({c}$4-1)*{c}$5")
line(10, "g50", "Net generation: P50", "MWh", INT, each=gen("cf50"), total=f"=SUM({C1}10:{CN}10)")
line(11, "g90", "Net generation: P90", "MWh", INT, each=gen("cf90"), total=f"=SUM({C1}11:{CN}11)")
line(12, "tar", "PPA tariff", "$/MWh", '#,##0.00', each=lambda c, p: f"={I['tariff']}*(1+{I['tesc']})^({c}$4-1)*{c}$5")
line(13, "r50", "Revenue: P50", "$mm", MM, each=lambda c, p: f"={c}10*{c}12/10^6", total=f"=SUM({C1}13:{CN}13)")
line(14, "r90", "Revenue: P90", "$mm", MM, each=lambda c, p: f"={c}11*{c}12/10^6", total=f"=SUM({C1}14:{CN}14)")
line(15, "opx", "Operating costs", "$mm", MM, each=lambda c, p: f"=-{I['cap']}*1000*{I['opex']}*(1+{I['oesc']})^({c}$4-1)/10^6*{c}$5", total=f"=SUM({C1}15:{CN}15)")
line(16, "e50", "EBITDA: P50", "$mm", MM, each=lambda c, p: f"={c}13+{c}15", total=f"=SUM({C1}16:{CN}16)")
line(17, "e90", "EBITDA: P90", "$mm", MM, each=lambda c, p: f"={c}14+{c}15", total=f"=SUM({C1}17:{CN}17)")
line(18, "dep", "Tax depreciation", "$mm", MM, each=lambda c, p: f"=IF({c}$4<={I['dep']},{I['cost']}/{I['dep']},0)*{c}$5", total=f"=SUM({C1}18:{CN}18)")
line(19, "t50", "Cash tax: P50", "$mm", MM, each=lambda c, p: f"=-MAX(0,{c}16-{c}18)*{I['tax']}", total=f"=SUM({C1}19:{CN}19)")
line(20, "t90", "Cash tax: P90", "$mm", MM, each=lambda c, p: f"=-MAX(0,{c}17-{c}18)*{I['tax']}", total=f"=SUM({C1}20:{CN}20)")
line(21, "c50", "CFADS: P50", "$mm", MM, each=lambda c, p: f"={c}16+{c}19", total=f"=SUM({C1}21:{CN}21)", bold=True)
line(22, "c90", "CFADS: P90", "$mm", MM, each=lambda c, p: f"={c}17+{c}20", total=f"=SUM({C1}22:{CN}22)", bold=True)

section(24, "Tranche A: Senior debt sizing (sculpted)")
line(25, "sa50", "Max DS @ P50 target DSCR", "$mm", MM, each=lambda c, p: f"={c}21/{I['s_d50']}*{c}$6")
line(26, "sa90", "Max DS @ P90 target DSCR", "$mm", MM, each=lambda c, p: f"={c}22/{I['s_d90']}*{c}$6")
line(27, "stgt", "Target senior DS (lower of)", "$mm", MM, each=lambda c, p: f"=MIN({c}25,{c}26)", total=f"=SUM({C1}27:{CN}27)")
line(28, "spv", "DSCR-based capacity (PV of target DS @ senior rate)", "$mm", MM, total=f"=NPV({I['s_rate']},{C1}27:{CN}27)")
line(29, "scap", "Gearing-based capacity", "$mm", MM, total=f"={I['s_gear']}*{I['cost']}")
line(30, "samt", "Senior debt amount", "$mm", MM, total="=MIN(C28,C29)", bold=True)
line(31, "sbind", "Binding constraint", "", "@", total='=IF(C29<C28,"Gearing","DSCR")')
line(32, "sscale", "DS scaling factor", "x", '0.0000', total="=IF(C28>0,C30/C28,0)")
line(33, "sds", "Senior debt service (sculpted)", "$mm", MM, each=lambda c, p: f"={c}27*$C$32", total=f"=SUM({C1}33:{CN}33)", bold=True)

section(35, "Tranche B: Mezzanine debt sizing (residual on combined DSCR)")
line(36, "ma50", "Max combined DS @ P50 floor", "$mm", MM, each=lambda c, p: f"={c}21/{I['m_d50']}*{c}$7")
line(37, "ma90", "Max combined DS @ P90 floor", "$mm", MM, each=lambda c, p: f"={c}22/{I['m_d90']}*{c}$7")
line(38, "mtgt", "Target mezz DS (combined less senior)", "$mm", MM, each=lambda c, p: f"=MAX(0,MIN({c}36,{c}37)-{c}33)*{c}$7", total=f"=SUM({C1}38:{CN}38)")
line(39, "mpv", "DSCR-based capacity (PV of target DS @ mezz rate)", "$mm", MM, total=f"=NPV({I['m_rate']},{C1}38:{CN}38)")
line(40, "mcap", "Gearing-based capacity (total cap less senior)", "$mm", MM, total=f"=MAX(0,{I['m_gear']}*{I['cost']}-C30)")
line(41, "mamt", "Mezzanine debt amount", "$mm", MM, total="=MIN(C39,C40)", bold=True)
line(42, "mbind", "Binding constraint", "", "@", total='=IF(C40<C39,"Gearing","DSCR")')
line(43, "mscale", "DS scaling factor", "x", '0.0000', total="=IF(C39>0,C41/C39,0)")
line(44, "mds", "Mezz debt service (sculpted)", "$mm", MM, each=lambda c, p: f"={c}38*$C$43", total=f"=SUM({C1}44:{CN}44)", bold=True)

section(46, "Senior debt schedule")
line(47, "sob", "Opening balance", "$mm", MM, first="=0", each=lambda c, p: f"={p}51")
line(48, "sdd", "Drawdown", "$mm", MM, first="=C30", each=lambda c, p: "=0", total=f"=SUM({FC}48:{CN}48)")
line(49, "sint", "Interest", "$mm", MM, first="=0", each=lambda c, p: f"={c}47*{I['s_rate']}", total=f"=SUM({FC}49:{CN}49)")
line(50, "sprin", "Principal repayment", "$mm", MM, first="=0", each=lambda c, p: f"={c}33-{c}49", total=f"=SUM({FC}50:{CN}50)")
line(51, "scb", "Closing balance", "$mm", MM, first=f"={FC}47+{FC}48-{FC}50", each=lambda c, p: f"={c}47+{c}48-{c}50", bold=True)
line(52, "sdscr", "Senior DSCR: P50", "x", X, each=lambda c, p: f'=IF({c}33>0.0001,{c}21/{c}33,"")')
line(53, "sdscr90", "Senior DSCR: P90", "x", X, each=lambda c, p: f'=IF({c}33>0.0001,{c}22/{c}33,"")')

section(55, "Mezzanine debt schedule")
line(56, "mob", "Opening balance", "$mm", MM, first="=0", each=lambda c, p: f"={p}60")
line(57, "mdd", "Drawdown", "$mm", MM, first="=C41", each=lambda c, p: "=0", total=f"=SUM({FC}57:{CN}57)")
line(58, "mint", "Interest", "$mm", MM, first="=0", each=lambda c, p: f"={c}56*{I['m_rate']}", total=f"=SUM({FC}58:{CN}58)")
line(59, "mprin", "Principal repayment", "$mm", MM, first="=0", each=lambda c, p: f"={c}44-{c}58", total=f"=SUM({FC}59:{CN}59)")
line(60, "mcb", "Closing balance", "$mm", MM, first=f"={FC}56+{FC}57-{FC}59", each=lambda c, p: f"={c}56+{c}57-{c}59", bold=True)
line(61, "cdscr", "Combined DSCR: P50", "x", X, each=lambda c, p: f'=IF({c}33+{c}44>0.0001,{c}21/({c}33+{c}44),"")')
line(62, "cdscr90", "Combined DSCR: P90", "x", X, each=lambda c, p: f'=IF({c}33+{c}44>0.0001,{c}22/({c}33+{c}44),"")')

section(64, "Equity & returns (P50 base case)")
line(65, "pcf", "Project cash flow (unlevered, post-tax)", "$mm", MM, first=f"=-{I['cost']}", each=lambda c, p: f"={c}21", total=f"=SUM({FC}65:{CN}65)")
line(66, "eq", "Equity cash flow", "$mm", MM, first=f"=-({I['cost']}-C30-C41)", each=lambda c, p: f"={c}21-{c}33-{c}44", total=f"=SUM({FC}66:{CN}66)", bold=True)

m.column_dimensions["A"].width = 48
m.column_dimensions["B"].width = 8
m.column_dimensions["C"].width = 13
for ci in range(FIRST, LAST + 1):
    m.column_dimensions[L(ci)].width = 10
m.freeze_panes = "D5"

# ---------------------------------------------------------------- Summary
s = wb.create_sheet("Summary", 0)
s["A1"] = "Tranched Debt Sizing: Summary"
s["A1"].font = TITLE
s["A2"] = "150 MW solar PV, illustrative inputs. All figures in $mm unless stated. Change blue cells on the Inputs tab."
s["A2"].font = Font(name="Arial", size=9, italic=True)


def hdr(r, cols):
    for i, h in enumerate(cols, 1):
        c = s.cell(r, i, h)
        c.font = HDR
        c.fill = HDR_FILL


hdr(4, ["Sources of funds", "$mm", "% of total", "Rate", "Tenor (yrs)", "Sizing constraint"])
src = [
    ("Tranche A: Senior term loan", "=Model!C30", "=Inputs!B25", "=Inputs!B24", "=Model!C31"),
    ("Tranche B: Mezzanine loan", "=Model!C41", "=Inputs!B31", "=Inputs!B30", "=Model!C42"),
    ("Sponsor equity", "=B9-B5-B6", None, None, "Plug"),
]
for i, (lab, amt, rate, ten, bind) in enumerate(src):
    r = 5 + i
    s.cell(r, 1, lab).font = BLACK
    for col, val, fmt in ((2, amt, MM), (4, rate, PCT), (5, ten, INT), (6, bind, "@")):
        if val is not None:
            c = s.cell(r, col, val)
            c.number_format = fmt
            c.font = GREEN if str(val).startswith("=Model") or str(val).startswith("=Inputs") else BLACK
    c = s.cell(r, 3, f"=IF($B$9>0,B{r}/$B$9,0)")
    c.number_format = PCT
    c.font = BLACK
s["A8"] = "Total debt"
s["B8"] = "=B5+B6"
s["C8"] = "=IF($B$9>0,B8/$B$9,0)"
s["A9"] = "Total sources"
s["B9"] = "=Inputs!B19"
s["C9"] = "=SUM(C5:C7)"
for r in (8, 9):
    s.cell(r, 1).font = BOLD
    s.cell(r, 2).font = BOLD
    s.cell(r, 2).number_format = MM
    s.cell(r, 3).font = BOLD
    s.cell(r, 3).number_format = PCT
    for col in range(1, 7):
        s.cell(r, col).border = TOP
s["B9"].font = Font(name="Arial", size=10, bold=True, color="008000")

hdr(11, ["Uses of funds", "$mm", "% of total"])
s["A12"], s["B12"] = "EPC + development costs", "=Inputs!B17"
s["A13"], s["B13"] = "Financing fees + IDC", "=Inputs!B18"
s["A14"], s["B14"] = "Total uses", "=B12+B13"
for r in (12, 13, 14):
    s.cell(r, 1).font = BOLD if r == 14 else BLACK
    s.cell(r, 2).font = GREEN if r < 14 else BOLD
    s.cell(r, 2).number_format = MM
    c = s.cell(r, 3, f"=IF($B$14>0,B{r}/$B$14,0)")
    c.number_format = PCT
    c.font = BOLD if r == 14 else BLACK
for col in range(1, 4):
    s.cell(14, col).border = TOP

hdr(16, ["Key credit metrics (P50 unless stated)", "Value", "", "", "", ""])
yrs = f"Model!{C1}$4:{CN}$4"
metrics = [
    ("Senior: minimum DSCR (P50)", f"=MIN(Model!{C1}52:{CN}52)", X),
    ("Senior: average DSCR (P50)", f"=AVERAGE(Model!{C1}52:{CN}52)", X),
    ("Senior: minimum DSCR (P90)", f"=MIN(Model!{C1}53:{CN}53)", X),
    ("Senior: LLCR at close (P50)", f"=IF(Model!C30>0,SUMPRODUCT(Model!{C1}21:{CN}21,Model!{C1}6:{CN}6,1/(1+Inputs!B25)^{yrs})/Model!C30,0)", X),
    ("Senior: weighted average life (yrs)", f"=IF(Model!C30>0,SUMPRODUCT(Model!{C1}50:{CN}50,{yrs})/Model!C30,0)", '0.0'),
    ("Combined: minimum DSCR (P50)", f"=MIN(Model!{C1}61:{CN}61)", X),
    ("Combined: minimum DSCR (P90)", f"=MIN(Model!{C1}62:{CN}62)", X),
    ("Mezz: weighted average life (yrs)", f"=IF(Model!C41>0,SUMPRODUCT(Model!{C1}59:{CN}59,{yrs})/Model!C41,0)", '0.0'),
    ("Project IRR (unlevered, post-tax)", f"=IRR(Model!{FC}65:{CN}65)", PCT),
    ("Equity IRR (post-tax, P50)", f"=IRR(Model!{FC}66:{CN}66,0.15)", PCT),
    ("Equity multiple (MOIC)", f"=IF(Model!{FC}66<0,SUM(Model!{C1}66:{CN}66)/-Model!{FC}66,0)", '0.00"x"'),
]
for i, (lab, f, fmt) in enumerate(metrics):
    r = 17 + i
    s.cell(r, 1, lab).font = BLACK
    c = s.cell(r, 2, f)
    c.number_format = fmt
    c.font = BLACK

r0 = 17 + len(metrics) + 1
hdr(r0, ["Integrity checks", "Value", "Status", "", "", ""])
checks = [
    ("Senior balance at maturity", f"=INDEX(Model!{C1}51:{CN}51,Inputs!B24)"),
    ("Mezz balance at maturity", f"=INDEX(Model!{C1}60:{CN}60,Inputs!B30)"),
    ("Sources less uses", "=B9-B14"),
    ("Minimum senior principal (no negative amortization)", f"=MIN(Model!{C1}50:{CN}50)"),
    ("Minimum mezz principal (no negative amortization)", f"=MIN(Model!{C1}59:{CN}59)"),
]
for i, (lab, f) in enumerate(checks):
    r = r0 + 1 + i
    s.cell(r, 1, lab).font = BLACK
    c = s.cell(r, 2, f)
    c.number_format = '#,##0.000;(#,##0.000);"-"'
    c.font = BLACK
    test = f"=IF(B{r}>=-0.001,\"OK\",\"CHECK\")" if "principal" in lab else f"=IF(ABS(B{r})<0.001,\"OK\",\"CHECK\")"
    s.cell(r, 3, test).font = BOLD

notes_r = r0 + len(checks) + 2
s.cell(notes_r, 1, "Structuring notes").font = BOLD
notes = [
    "1. Senior debt service is sculpted to the lower of CFADS/1.30x (P50) and CFADS/1.20x (P90) over its tenor; debt = PV of that profile at the all-in rate, capped at max gearing.",
    "2. If gearing binds, the sculpted profile is scaled down pro rata so the loan still fully amortizes at maturity (actual DSCRs rise above target).",
    "3. Mezz sizes on combined DSCR floors (1.15x P50 / 1.05x P90) and takes the residual DS capacity after senior. If mezz tenor exceeds senior, most of its capacity sits in the tail and early years can turn into PIK accretion (flagged by the integrity check).",
    "4. Simplifications: annual periods, no construction-period drawdown profile, no DSRA/MRA, no cash sweep, no interest tax shield. Fees/IDC held fixed to stay non-circular.",
]
for i, n in enumerate(notes):
    s.cell(notes_r + 1 + i, 1, n).font = Font(name="Arial", size=9)

s.column_dimensions["A"].width = 50
for col in "BCDEF":
    s.column_dimensions[col].width = 14
s.column_dimensions["F"].width = 18

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False

import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved", OUT)
